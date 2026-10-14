import streamlit as st
from PyPDF2 import PdfReader
import docx
import pandas as pd
import json
import os
import hashlib
from llama_index.core import VectorStoreIndex, Document, Settings, StorageContext, load_index_from_storage
from llama_index.embeddings.ollama import OllamaEmbedding
from llama_index.llms.ollama import Ollama
from llama_index.core.node_parser import SentenceSplitter
from cache_manager import BotCacheManager

# Initialize cache manager
cache_manager = BotCacheManager()

# Configure LlamaIndex settings
Settings.llm = Ollama(model="mistral", request_timeout=1200.0)
Settings.embed_model = OllamaEmbedding(model_name="nomic-embed-text")

st.header("My Document Chatbot (LlamaIndex)")

def get_file_hash(file_content):
    """Generate a hash for the file content to detect changes"""
    return hashlib.md5(file_content.encode()).hexdigest()

def save_index_safe(index, file_hash):
    """Safely save LlamaIndex to disk"""
    try:
        # Create a cache directory for indexes
        cache_dir = "llamaindex_cache"
        if not os.path.exists(cache_dir):
            os.makedirs(cache_dir)
        
        # Save the index
        index_path = os.path.join(cache_dir, f"index_{file_hash}")
        index.storage_context.persist(persist_dir=index_path)
        
        # Cache the file hash for quick lookup
        cache_manager.cache_response(f"llamaindex_{file_hash}", {
            "index_path": index_path,
            "file_hash": file_hash,
            "status": "saved"
        })
        
        return True
    except Exception as e:
        st.warning(f"Could not cache index: {str(e)}")
        return False

def load_index_safe(file_hash):
    """Safely load LlamaIndex from disk"""
    try:
        # Check if we have cached info about this index
        cached_info = cache_manager.get_cached_response(f"llamaindex_{file_hash}")
        
        if cached_info and "index_path" in cached_info:
            index_path = cached_info["index_path"]
            
            # Check if the index directory exists
            if os.path.exists(index_path):
                # Load the storage context and index
                storage_context = StorageContext.from_defaults(persist_dir=index_path)
                index = load_index_from_storage(storage_context)
                return index
        
        return None
    except Exception as e:
        st.warning(f"Could not load cached index: {str(e)}")
        return None

def extract_text_from_file(file):
    """Extract text from various file formats"""
    text = ""
    file_type = file.type
    
    try:
        if file_type == "application/pdf":
            # PDF files
            pdf_reader = PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text()
                
        elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            # DOCX files
            doc = docx.Document(file)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
                
        elif file_type == "text/plain":
            # TXT files
            text = str(file.read(), "utf-8")
            
        elif file_type == "text/csv":
            # CSV files
            df = pd.read_csv(file)
            text = df.to_string()
            
        elif file_type in ["application/vnd.ms-excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"]:
            # Excel files
            df = pd.read_excel(file)
            text = df.to_string()
            
        elif file_type == "application/json":
            # JSON files
            json_data = json.load(file)
            text = json.dumps(json_data, indent=2)
            
        elif file_type == "text/markdown":
            # Markdown files
            text = str(file.read(), "utf-8")
            
        elif file_type in ["application/rtf", "text/rtf"]:
            # RTF files (basic support)
            text = str(file.read(), "utf-8")
            
        else:
            # Try to read as text for other formats
            try:
                text = str(file.read(), "utf-8")
            except:
                st.error(f"Unsupported file format: {file_type}")
                return None
                
    except Exception as e:
        st.error(f"Error reading file: {str(e)}")
        return None
        
    return text

with st.sidebar:
    st.title("Your Documents")
    
    # Enhanced file uploader with multiple formats
    file = st.file_uploader(
        "Upload any document and start asking questions", 
        type=[
            "pdf", "docx", "txt", "csv", "xlsx", "xls", 
            "json", "md", "rtf", "py", "js", "html", "xml", "log"
        ]
    )
    
    if file is not None:
        st.info(f"📄 File: {file.name}")
        st.info(f"📊 Size: {file.size / 1024:.1f} KB")
        st.info(f"🔧 Type: {file.type}")

# Extract the text and create index
if file is not None:
    with st.spinner(f"📄 Processing {file.name}..."):
        text = extract_text_from_file(file)
        
    if text is None:
        st.stop()
        
    if len(text.strip()) == 0:
        st.error("No text content found in the file!")
        st.stop()
    
    # Generate hash for the file content to create unique cache
    file_hash = get_file_hash(text)
    
    # Try to load cached index
    with st.spinner("🚀 Checking for cached index..."):
        index = load_index_safe(file_hash)
        
    if index is not None:
        st.success("✅ Loaded index from cache! Ready to answer questions.")
    else:
        # If no cached index, create new one
        with st.spinner("🧠 Creating document index..."):
            # Create Document object
            document = Document(text=text, metadata={"filename": file.name})
            
            # Configure node parser (equivalent to text splitter)
            node_parser = SentenceSplitter(
                chunk_size=600,
                chunk_overlap=100,
            )
            
            # Create index
            index = VectorStoreIndex.from_documents(
                [document],
                node_parser=node_parser,
                show_progress=True
            )
        
        # Cache the index for future use
        with st.spinner("💾 Caching index..."):
            if save_index_safe(index, file_hash):
                st.success("✅ Document processed and cached! Ask your questions below.")
            else:
                st.success("✅ Document ready! Ask your questions below.")

    # Create query engine
    query_engine = index.as_query_engine(
        similarity_top_k=5,
        response_mode="compact"
    )

    # Initialize session state for conversation history
    if "messages" not in st.session_state:
        st.session_state.messages = []

    # Display conversation history
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.write(message["content"])

    # Chat input
    if prompt := st.chat_input("Ask a question about your document"):
        # Add user message to chat history
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.write(prompt)

        # Generate response
        with st.chat_message("assistant"):
            with st.spinner("🤔 Generating answer..."):
                # Query the index
                response = query_engine.query(prompt)
                answer = str(response)
                st.write(answer)
                
                # Add assistant response to chat history
                st.session_state.messages.append({"role": "assistant", "content": answer})