import json
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Union

import docx
import pandas as pd
from llama_index.core import Document
from pypdf import PdfReader

PathLike = Union[str, Path]


def _metadata(filename: str, **extra) -> dict:
    meta = {"filename": filename}
    meta.update(extra)
    return meta


def documents_from_pdf(file_obj: BinaryIO, filename: str) -> list[Document]:
    reader = PdfReader(file_obj)
    documents: list[Document] = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        documents.append(
            Document(
                text=text,
                metadata=_metadata(
                    filename,
                    page_label=str(page_num),
                    page_number=page_num,
                ),
            )
        )
    return documents


def documents_from_docx(file_obj: BinaryIO, filename: str) -> list[Document]:
    doc = docx.Document(file_obj)
    text = "\n".join(p.text for p in doc.paragraphs)
    if not text.strip():
        return []
    return [Document(text=text, metadata=_metadata(filename))]


def documents_from_text(content: str, filename: str) -> list[Document]:
    if not content.strip():
        return []
    return [Document(text=content, metadata=_metadata(filename))]


def documents_from_csv(file_obj: BinaryIO, filename: str) -> list[Document]:
    df = pd.read_csv(file_obj)
    text = df.to_string()
    return documents_from_text(text, filename)


def documents_from_excel(file_obj: BinaryIO, filename: str) -> list[Document]:
    df = pd.read_excel(file_obj)
    text = df.to_string()
    return documents_from_text(text, filename)


def documents_from_json(file_obj: BinaryIO, filename: str) -> list[Document]:
    data = json.load(file_obj)
    text = json.dumps(data, indent=2)
    return documents_from_text(text, filename)


def extract_documents_from_upload(file, filename: str | None = None) -> list[Document]:
    """Extract LlamaIndex Documents from a Streamlit UploadedFile or file-like object."""
    name = filename or getattr(file, "name", "upload")
    file_type = getattr(file, "type", "") or ""

    try:
        if file_type == "application/pdf" or name.lower().endswith(".pdf"):
            return documents_from_pdf(file, name)

        if (
            file_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or name.lower().endswith(".docx")
        ):
            return documents_from_docx(file, name)

        if file_type == "text/csv" or name.lower().endswith(".csv"):
            return documents_from_csv(file, name)

        if file_type in (
            "application/vnd.ms-excel",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ) or name.lower().endswith((".xlsx", ".xls")):
            return documents_from_excel(file, name)

        if file_type == "application/json" or name.lower().endswith(".json"):
            return documents_from_json(file, name)

        if file_type in ("text/plain", "text/markdown", "application/rtf", "text/rtf") or name.lower().endswith(
            (".txt", ".md", ".rtf", ".py", ".js", ".html", ".xml", ".log")
        ):
            raw = file.read()
            if isinstance(raw, bytes):
                content = raw.decode("utf-8", errors="replace")
            else:
                content = str(raw)
            return documents_from_text(content, name)

        raw = file.read()
        content = raw.decode("utf-8", errors="replace") if isinstance(raw, bytes) else str(raw)
        return documents_from_text(content, name)
    except Exception:
        raise


def load_documents_from_path(path: PathLike) -> list[Document]:
    """Load documents from a file on disk (for evals and scripts)."""
    path = Path(path)
    filename = path.name
    suffix = path.suffix.lower()

    with path.open("rb") as f:
        if suffix == ".pdf":
            return documents_from_pdf(f, filename)
        if suffix == ".docx":
            return documents_from_docx(f, filename)
        if suffix == ".csv":
            return documents_from_csv(f, filename)
        if suffix in (".xlsx", ".xls"):
            return documents_from_excel(f, filename)
        if suffix == ".json":
            return documents_from_json(f, filename)

    content = path.read_text(encoding="utf-8", errors="replace")
    return documents_from_text(content, filename)


def load_documents_from_directory(directory: PathLike) -> list[Document]:
    directory = Path(directory)
    docs: list[Document] = []
    for path in sorted(directory.iterdir()):
        if path.is_file() and not path.name.startswith("."):
            docs.extend(load_documents_from_path(path))
    return docs


def content_fingerprint(documents: list[Document]) -> str:
    import hashlib

    joined = "\n".join(d.text for d in documents)
    return hashlib.md5(joined.encode()).hexdigest()


def read_upload_as_bytes(file) -> BytesIO:
    """Reset-friendly read for Streamlit uploads."""
    data = file.getvalue() if hasattr(file, "getvalue") else file.read()
    return BytesIO(data)
